import sys
from pptx import Presentation
from pptx.util import Emu

def emu_to_in(v):
    return round(Emu(v).inches, 2) if v is not None else None

def shape_fill(shape):
    try:
        fill = shape.fill
        if fill.type is None:
            return "None"
        if hasattr(fill, "fore_color"):
            rgb = fill.fore_color.rgb
            return str(rgb) if rgb else "None"
    except Exception:
        pass
    return "None"

def shape_text(shape):
    try:
        if shape.has_text_frame:
            return shape.text_frame.text.replace("\n", " | ")[:60]
    except Exception:
        pass
    return ""

path = sys.argv[1]
prs = Presentation(path)
for i, slide in enumerate(prs.slides, 1):
    print(f"--- SLIDE {i} ---")
    for shape in slide.shapes:
        x, y = emu_to_in(shape.left), emu_to_in(shape.top)
        w, h = emu_to_in(shape.width), emu_to_in(shape.height)
        stype = str(shape.shape_type).split(" ")[0] if shape.shape_type else "?"
        fill = shape_fill(shape) if shape.shape_type != 13 else "None"
        text = shape_text(shape)
        print(f"  [{stype}] pos=({x},{y}) size=({w}x{h}) fill={fill} txt='{text}'")
